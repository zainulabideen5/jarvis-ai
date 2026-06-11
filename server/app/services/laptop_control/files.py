"""File operations — find, read, open, edit, move, delete."""

from __future__ import annotations

import os
import re
import shutil
import string
import subprocess
import time
from pathlib import Path

from app.core.logging import get_logger

log = get_logger(__name__)

# Common user folders
HOME = Path.home()

# Folders to skip during deep scan (system + dev junk)
EXCLUDED_DIRS = {
    "windows", "program files", "program files (x86)", "programdata",
    "appdata", "$recycle.bin", "system volume information", "msocache",
    "node_modules", ".git", "__pycache__", ".venv", "venv", "env",
    "dist", "build", ".cache", ".npm", ".vscode", ".idea",
    ".next", ".nuxt", "target", "out", "bin", "obj",
}


def _get_all_drives() -> list[Path]:
    """Return list of all available drives."""
    drives = []
    for letter in string.ascii_uppercase:
        d = Path(f"{letter}:\\")
        try:
            if d.exists():
                drives.append(d)
        except OSError:
            continue
    return drives


def _is_excluded(path: Path) -> bool:
    """Check if any path part is in excluded list."""
    try:
        return any(p.lower() in EXCLUDED_DIRS for p in path.parts)
    except Exception:
        return False


def _windows_search_index(query: str, max_results: int = 30) -> list[dict]:
    """Use Windows Search Index for fast full-system search."""
    if not query or len(query) < 2:
        return []
    safe_query = query.replace("'", "''")
    ps_script = (
        "$ErrorActionPreference='SilentlyContinue';"
        "$conn=New-Object -ComObject ADODB.Connection;"
        "$conn.Open(\"Provider=Search.CollatorDSO;Extended Properties='Application=Windows';\");"
        f"$rs=$conn.Execute(\"SELECT TOP {max_results} System.ItemPathDisplay,System.ItemType,System.Size,System.DateModified "
        f"FROM SystemIndex WHERE System.FileName LIKE '%{safe_query}%' OR System.ItemNameDisplay LIKE '%{safe_query}%'\");"
        "while(-not $rs.EOF){"
        "$path=$rs.Fields.Item(0).Value;"
        "$kind=$rs.Fields.Item(1).Value;"
        "$size=$rs.Fields.Item(2).Value;"
        "$mod=$rs.Fields.Item(3).Value;"
        "Write-Output (\"$path|$kind|$size|$mod\");"
        "$rs.MoveNext()};"
        "$rs.Close();$conn.Close()"
    )
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps_script],
            capture_output=True, text=True, timeout=15,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        out = result.stdout or ""
        items = []
        for line in out.splitlines():
            line = line.strip()
            if not line or "|" not in line:
                continue
            parts = line.split("|")
            if len(parts) < 2:
                continue
            full_path = parts[0].strip()
            kind = parts[1].strip().lower()
            try:
                p = Path(full_path)
                if not p.exists():
                    continue
                is_folder = p.is_dir() or "directory" in kind or "folder" in kind
                stat = p.stat()
                items.append({
                    "name": p.name,
                    "path": str(p),
                    "size_kb": 0 if is_folder else round(stat.st_size / 1024, 1),
                    "modified": stat.st_mtime,
                    "folder": str(p.parent.name),
                    "is_folder": is_folder,
                })
            except (OSError, ValueError):
                continue
        return items
    except (subprocess.TimeoutExpired, Exception) as e:
        log.debug("windows_search_index_failed", error=str(e))
        return []

def _get_known_folder(csidl: int) -> Path | None:
    """Get a Windows special folder by CSIDL — handles OneDrive moves correctly."""
    if os.name != "nt":
        return None
    try:
        import ctypes
        from ctypes import wintypes
        buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
        ctypes.windll.shell32.SHGetFolderPathW(None, csidl, None, 0, buf)
        if buf.value:
            p = Path(buf.value)
            return p if p.exists() else None
    except Exception:
        pass
    return None


def _build_special_folders() -> dict:
    """Resolve real Windows folder paths — handles OneDrive."""
    desktop = _get_known_folder(0) or (HOME / "Desktop")        # CSIDL_DESKTOP
    documents = _get_known_folder(5) or (HOME / "Documents")    # CSIDL_PERSONAL
    pictures = _get_known_folder(39) or (HOME / "Pictures")     # CSIDL_MYPICTURES
    music = _get_known_folder(13) or (HOME / "Music")           # CSIDL_MYMUSIC
    videos = _get_known_folder(14) or (HOME / "Videos")         # CSIDL_MYVIDEO
    downloads = HOME / "Downloads"
    if not downloads.exists():
        od = HOME / "OneDrive" / "Downloads"
        if od.exists():
            downloads = od

    return {
        "desktop": desktop,
        "documents": documents,
        "docs": documents,
        "downloads": downloads,
        "download": downloads,
        "pictures": pictures,
        "images": pictures,
        "photos": pictures,
        "videos": videos,
        "video": videos,
        "music": music,
    }


SPECIAL_FOLDERS = _build_special_folders()

# Search paths: real folders + legacy variants (some users have both)
SEARCH_PATHS = []
_seen_paths = set()
for _key in ("desktop", "documents", "downloads", "pictures", "videos", "music"):
    _p = SPECIAL_FOLDERS.get(_key)
    if _p and str(_p) not in _seen_paths and _p.exists():
        SEARCH_PATHS.append(_p)
        _seen_paths.add(str(_p))
for _legacy in (HOME / "Desktop", HOME / "Documents", HOME / "Downloads"):
    if _legacy.exists() and str(_legacy) not in _seen_paths:
        SEARCH_PATHS.append(_legacy)
        _seen_paths.add(str(_legacy))

log.info("special_folders_resolved", desktop=str(SPECIAL_FOLDERS["desktop"]), documents=str(SPECIAL_FOLDERS["documents"]))

MAX_RESULTS = 20


def _scan_paths(bases: list[Path], query_lower: str, max_results: int,
                include_folders: bool = True, deep: bool = True,
                time_budget_sec: float = 8.0) -> list[dict]:
    """Manual filesystem scan with smart excludes and time budget."""
    results = []
    deadline = time.monotonic() + time_budget_sec

    for base in bases:
        if not base.exists():
            continue
        try:
            iterator = base.rglob("*") if deep else _shallow_iter(base, max_depth=4)
            for path in iterator:
                if time.monotonic() > deadline:
                    return results
                if len(results) >= max_results * 3:
                    break
                if _is_excluded(path):
                    continue
                try:
                    is_folder = path.is_dir()
                    if not include_folders and is_folder:
                        continue
                    if not is_folder and not path.is_file():
                        continue
                    if query_lower in path.name.lower():
                        stat = path.stat()
                        results.append({
                            "name": path.name,
                            "path": str(path),
                            "size_kb": 0 if is_folder else round(stat.st_size / 1024, 1),
                            "modified": stat.st_mtime,
                            "folder": str(path.parent.name),
                            "is_folder": is_folder,
                        })
                except (OSError, ValueError):
                    continue
        except (PermissionError, OSError):
            continue
    return results


def _shallow_iter(base: Path, max_depth: int = 4):
    """Iterate paths up to max_depth, skipping excluded folders early."""
    def _walk(current: Path, depth: int):
        if depth > max_depth or _is_excluded(current):
            return
        try:
            for entry in current.iterdir():
                yield entry
                if entry.is_dir() and not _is_excluded(entry):
                    yield from _walk(entry, depth + 1)
        except (PermissionError, OSError):
            return
    yield from _walk(base, 0)


def _resolve_location(location: str | None) -> Path:
    """Resolve a folder name (case-insensitive) to a real Path."""
    if not location:
        return HOME / "Desktop"
    key = location.lower().strip()
    if key in SPECIAL_FOLDERS:
        return SPECIAL_FOLDERS[key]
    # Maybe it's a full path
    try:
        p = Path(location).expanduser()
        if p.is_absolute() and p.exists():
            return p
    except Exception:
        pass
    return HOME / "Desktop"


class FileOperations:
    """Safe file operations."""

    @staticmethod
    def _extract_search_terms(query: str) -> list[str]:
        """Extract meaningful search terms from natural-language query."""
        # Strip common Urdu/Hindi/English filler words
        STOPWORDS = {
            "ka", "ki", "ke", "ko", "wala", "wali", "wale", "naam", "naame", "vala", "vali",
            "the", "a", "an", "of", "is", "in", "at", "on", "for", "and", "or",
            "file", "files", "folder", "folders", "wala", "type",
            "kholo", "open", "khol", "padho", "find", "dhondo", "search", "dikhao",
            "yr", "yaar", "bhai", "please", "pls",
        }
        words = re.findall(r"[\w]+", query.lower())
        terms = [w for w in words if w and w not in STOPWORDS and len(w) >= 2]
        return terms or [query.lower().strip()]

    @staticmethod
    def find_files(query: str, location: str | None = None, max_results: int = MAX_RESULTS,
                   include_folders: bool = True, prefer_kind: str = "auto") -> list[dict]:
        """Find files AND folders matching query.
        prefer_kind: 'file' | 'folder' | 'auto' — sorts matching kind first."""
        query_clean = (query or "").strip()
        if not query_clean:
            return []
        query_lower = query_clean.lower()

        def _do_search(q: str) -> list[dict]:
            if location:
                base = _resolve_location(location)
                return _scan_paths([base] if base.exists() else SEARCH_PATHS,
                                   q, max_results, include_folders)
            r = _windows_search_index(q, max_results=max_results * 2)
            if include_folders is False:
                r = [x for x in r if not x.get("is_folder")]
            if not r:
                paths = list(SEARCH_PATHS) + _get_all_drives()
                r = _scan_paths(paths, q, max_results, include_folders, deep=False)
            return r

        # 1) Try full query
        results = _do_search(query_lower)

        # 2) If nothing found, try without extension (e.g., "zain.text" -> "zain")
        if not results and "." in query_clean:
            stem = Path(query_clean).stem
            if stem and stem != query_clean:
                results = _do_search(stem.lower())

        # 3) If still nothing, split query into terms — longest/rarest first
        if not results and (" " in query_clean or "." in query_clean or "_" in query_clean or "-" in query_clean):
            terms = FileOperations._extract_search_terms(query_clean)
            terms = sorted(set(terms), key=len, reverse=True)
            for term in terms:
                if len(term) < 3:
                    continue
                partial = _do_search(term)
                if partial:
                    results = partial
                    break

        # Sort: exact name match first, term match, kind preference, then recency
        terms = FileOperations._extract_search_terms(query_clean)

        def _score(r):
            name_lower = r["name"].lower()
            stem = Path(r["name"]).stem.lower()
            is_folder = r.get("is_folder", False)
            exact = name_lower == query_lower
            stem_exact = stem == query_lower
            term_stem = any(stem == t for t in terms)

            # Kind preference
            if prefer_kind == "file":
                kind_score = 0 if is_folder else -1  # files first
            elif prefer_kind == "folder":
                kind_score = -1 if is_folder else 0  # folders first
            else:
                kind_score = -1 if is_folder else 0  # default: folders first

            return (
                -1 if exact else 0,
                -1 if stem_exact else 0,
                -1 if term_stem else 0,
                kind_score,
                -r.get("modified", 0),
            )
        results.sort(key=_score)
        return results[:max_results]

    @staticmethod
    def open_folder(name: str, location: str | None = None) -> tuple[bool, str]:
        """Find and open a folder in Explorer."""
        if not name or not name.strip():
            return False, "Folder ka naam dena hoga"

        name_clean = name.strip()
        name_lower = name_clean.lower()

        # If name itself is a special folder (Desktop, Documents, Downloads, etc.) — open it directly
        if name_lower in SPECIAL_FOLDERS:
            target = SPECIAL_FOLDERS[name_lower]
            if target.exists():
                return FileOperations._shell_open(target)

        # If location is given AND name matches the location → just open location
        if location and location.lower() == name_lower and location.lower() in SPECIAL_FOLDERS:
            target = SPECIAL_FOLDERS[location.lower()]
            if target.exists():
                return FileOperations._shell_open(target)

        # Try direct absolute path
        try:
            p = Path(name_clean).expanduser()
            if p.is_absolute() and p.exists() and p.is_dir():
                return FileOperations._shell_open(p)
        except Exception:
            pass

        # Try inside specified location
        if location:
            base = _resolve_location(location)
            candidate = base / name_clean
            if candidate.exists() and candidate.is_dir():
                return FileOperations._shell_open(candidate)

        # Full search
        results = FileOperations.find_files(name_clean, location=location, include_folders=True)
        folders = [r for r in results if r.get("is_folder")]
        if not folders:
            return False, f"Folder '{name}' nahi mila — sure naam sahi hai?"

        target = Path(folders[0]["path"])
        if not target.exists():
            return False, f"Path mil gaya tha but ab nahi hai: {target}"
        return FileOperations._shell_open(target)

    @staticmethod
    def _shell_open(target: Path) -> tuple[bool, str]:
        """Open a path using Windows Shell API (most reliable from any context)."""
        target_str = str(target.resolve())
        kind = "Folder" if target.is_dir() else "File"
        icon = "📁" if target.is_dir() else "📄"
        success_msg = f"{kind} khol di:\n{icon} {target.name}\n📍 {target_str}"
        errors = []

        # Method 1: Windows ShellExecuteW — same API Explorer uses
        if os.name == "nt":
            try:
                import ctypes
                ret = ctypes.windll.shell32.ShellExecuteW(
                    None, "open", target_str, None, None, 1
                )
                if ret > 32:
                    return True, success_msg
                errors.append(f"ShellExecute returned {ret}")
            except Exception as e:
                errors.append(f"ShellExecute: {e}")

        # Method 2: explorer.exe for folders
        if target.is_dir():
            try:
                subprocess.Popen(
                    f'explorer "{target_str}"',
                    shell=True,
                    creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0,
                )
                return True, success_msg
            except Exception as e:
                errors.append(f"explorer: {e}")

        # Method 3: start command via cmd
        try:
            subprocess.Popen(
                f'start "" "{target_str}"',
                shell=True,
                creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0,
            )
            return True, success_msg
        except Exception as e:
            errors.append(f"start: {e}")

        # Method 4: os.startfile as last resort
        try:
            os.startfile(target_str)
            return True, success_msg
        except Exception as e:
            errors.append(f"startfile: {e}")

        log.warning("shell_open_all_methods_failed", path=target_str, errors=errors)
        return False, f"Open nahi hua: {'; '.join(errors[:2])}"

    @staticmethod
    def open_file(path: str) -> tuple[bool, str]:
        """Open a file or folder with default app/Explorer."""
        if not path or not str(path).strip():
            return False, "File ya folder ka naam dena hoga"

        # Try as direct path first
        try:
            p = Path(path).expanduser()
            if p.is_absolute() and p.exists():
                return FileOperations._shell_open(p)
        except Exception:
            pass

        # Search system
        matches = FileOperations.find_files(path, include_folders=True, max_results=1)
        if not matches:
            return False, f"Nahi mila: {path}"
        target = Path(matches[0]["path"])
        if not target.exists():
            return False, f"Path mil gaya tha but ab nahi hai: {target}"
        return FileOperations._shell_open(target)

    @staticmethod
    def read_file(path: str, max_chars: int = 5000) -> tuple[bool, str]:
        """Read file content (text, PDF, Word, Excel)."""
        try:
            p = Path(path).expanduser()
            if not p.exists():
                matches = FileOperations.find_files(path, max_results=1)
                if matches:
                    p = Path(matches[0]["path"])
                else:
                    return False, f"File nahi mili: {path}"

            ext = p.suffix.lower()

            if ext in (".txt", ".md", ".log", ".csv", ".json", ".py", ".js", ".html"):
                content = p.read_text(encoding="utf-8", errors="ignore")
                return True, content[:max_chars]

            if ext == ".pdf":
                try:
                    from PyPDF2 import PdfReader
                    reader = PdfReader(str(p))
                    text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
                    return True, text[:max_chars]
                except Exception as e:
                    return False, f"PDF read error: {e}"

            if ext == ".docx":
                try:
                    from docx import Document
                    doc = Document(str(p))
                    text = "\n".join(p.text for p in doc.paragraphs)
                    return True, text[:max_chars]
                except Exception as e:
                    return False, f"Word read error: {e}"

            if ext == ".xlsx":
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(str(p), data_only=True)
                    lines = []
                    for sheet in wb.sheetnames[:3]:
                        ws = wb[sheet]
                        lines.append(f"=== Sheet: {sheet} ===")
                        for row in ws.iter_rows(max_row=50, values_only=True):
                            lines.append(" | ".join(str(c) if c is not None else "" for c in row))
                    return True, "\n".join(lines)[:max_chars]
                except Exception as e:
                    return False, f"Excel read error: {e}"

            return False, f"Ye file type support nahi: {ext}"

        except Exception as e:
            return False, f"Error: {e}"

    @staticmethod
    def create_folder(name: str, location: str = "Desktop") -> tuple[bool, str]:
        """Create a folder."""
        if not name or not name.strip():
            return False, "Folder ka naam dena hoga"
        try:
            base = _resolve_location(location)
            folder_path = base / name.strip()
            if folder_path.exists():
                return True, f"Folder pehle se hai: {folder_path}"
            folder_path.mkdir(parents=True, exist_ok=True)
            return True, f"Folder bana di: {folder_path}"
        except Exception as e:
            return False, f"Folder bana nahi: {e}"

    @staticmethod
    def create_text_file(name: str, content: str, folder: str = "Desktop") -> tuple[bool, str]:
        """Create a text file."""
        if not name or not name.strip():
            return False, "File ka naam dena hoga"
        try:
            base = _resolve_location(folder)
            if not name.endswith(".txt") and "." not in name:
                name += ".txt"
            p = base / name.strip()
            p.write_text(content, encoding="utf-8")
            return True, f"File bana di: {p}"
        except Exception as e:
            return False, f"Create failed: {e}"

    @staticmethod
    def create_word_file(name: str, content: str, folder: str = "Desktop") -> tuple[bool, str]:
        """Create a Word document."""
        if not name or not name.strip():
            return False, "File ka naam dena hoga"
        try:
            from docx import Document
            base = _resolve_location(folder)
            if not name.endswith(".docx"):
                name += ".docx"
            doc = Document()
            for line in (content or "").split("\n"):
                doc.add_paragraph(line)
            p = base / name.strip()
            doc.save(str(p))
            return True, f"Word doc bana: {p}"
        except Exception as e:
            return False, f"Create failed: {e}"

    @staticmethod
    def add_excel_row(file: str, data: list, sheet: str | None = None) -> tuple[bool, str]:
        """Add a row to Excel file."""
        try:
            from openpyxl import load_workbook
            p = Path(file).expanduser()
            if not p.exists():
                matches = FileOperations.find_files(file, max_results=1)
                if matches:
                    p = Path(matches[0]["path"])
                else:
                    return False, f"Excel file nahi mili: {file}"

            wb = load_workbook(str(p))
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
            ws.append([str(d) if d is not None else "" for d in data])
            wb.save(str(p))
            return True, f"Row add ho gayi: {p.name}"
        except Exception as e:
            return False, f"Excel edit failed: {e}"

    @staticmethod
    def move_file(source: str, dest: str) -> tuple[bool, str]:
        """Move a file."""
        try:
            src = Path(source).expanduser()
            if not src.exists():
                return False, f"Source file nahi mili: {source}"

            dst = Path(dest).expanduser()
            if dst.is_dir():
                dst = dst / src.name

            shutil.move(str(src), str(dst))
            return True, f"Move ho gaya: {src.name} → {dst.parent.name}"
        except Exception as e:
            return False, f"Move failed: {e}"

    @staticmethod
    def delete_file(path: str, kind: str = "auto") -> tuple[bool, str]:
        """Delete a file or folder. kind: 'file' | 'folder' | 'auto'."""
        if not path or not path.strip():
            return False, "Path dena hoga"

        # Auto-detect from path string
        path_lower = path.lower()
        if kind == "auto":
            if "folder" in path_lower or "directory" in path_lower:
                kind = "folder"
            elif "file" in path_lower or "." in Path(path).suffix:
                kind = "file"

        target = None
        # Try absolute path first
        try:
            p = Path(path).expanduser()
            if p.is_absolute() and p.exists():
                target = p
        except Exception:
            pass

        # Search if not direct path
        if target is None:
            prefer = kind if kind in ("file", "folder") else "file"  # default: prefer file for delete
            include_folders = True
            matches = FileOperations.find_files(path, include_folders=include_folders,
                                                 max_results=5, prefer_kind=prefer)
            if not matches:
                return False, f"Nahi mila: {path}"

            # If kind is explicit, filter
            if kind == "file":
                files_only = [m for m in matches if not m.get("is_folder")]
                if files_only:
                    matches = files_only
            elif kind == "folder":
                folders_only = [m for m in matches if m.get("is_folder")]
                if folders_only:
                    matches = folders_only

            target = Path(matches[0]["path"])

        if not target.exists():
            return False, f"File/folder nahi mili: {target}"

        target_str = str(target.resolve())

        # Safety: refuse to delete anywhere outside the user's profile or removable
        # drives. Even though consent is granted, the AI must never `rm -rf C:\Windows`
        # because an LLM hallucinated a path. Allowed roots:
        #   - HOME and everything under it (Desktop/Documents/Downloads/OneDrive/etc.)
        #   - Any drive other than C: (D:, E:, USB) — user's project/scratch data
        # Blocked: C:\Windows, C:\Program Files, C:\ProgramData, C:\Users\Default, etc.
        target_low = target_str.lower()
        home_low = str(Path.home().resolve()).lower()
        BLOCKED_PREFIXES = (
            "c:\\windows",
            "c:\\program files",
            "c:\\program files (x86)",
            "c:\\programdata",
            "c:\\users\\default",
            "c:\\users\\public",
            "c:\\$recycle.bin",
            "c:\\system volume information",
        )
        in_home = target_low.startswith(home_low)
        on_other_drive = len(target_low) >= 2 and target_low[1] == ":" and not target_low.startswith("c:")
        if any(target_low.startswith(p) for p in BLOCKED_PREFIXES) or (not in_home and not on_other_drive):
            log.warning("delete_blocked_system_path", path=target_str)
            return False, (
                f"🛡️ Safety block: '{target_str}' system folder/path hai — delete nahi karunga. "
                f"Sirf tumhare profile aur removable drives ke andar delete allowed hai."
            )

        try:
            if target.is_dir():
                shutil.rmtree(target_str)
                return True, f"🗑️ Folder delete ho gaya:\n📁 {target.name}\n📍 {target_str}"
            else:
                target.unlink()
                return True, f"🗑️ File delete ho gayi:\n📄 {target.name}\n📍 {target_str}"
        except PermissionError as e:
            return False, f"Permission denied — file/folder use mein hogi: {e}"
        except Exception as e:
            return False, f"Delete failed: {e}"
